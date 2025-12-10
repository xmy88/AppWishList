import os
import re
import threading
import time
import datetime as dt
import logging
import traceback
import warnings
import sqlite3
import configparser
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog

import ttkbootstrap as tb
from ttkbootstrap.constants import *

from docx import Document
from docx.enum.text import WD_UNDERLINE, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from dateutil.relativedelta import relativedelta
from datetime import timedelta

warnings.filterwarnings("ignore", category=UserWarning, module='openpyxl')

logging.basicConfig(
    filename='contract_tool.log',
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)


class FreeMaintenanceContractApp:
    """原始免保合同制作功能，保持原有逻辑。"""

    def __init__(self, parent):
        self.parent = parent
        self.config = configparser.ConfigParser()
        self.config_file = 'config.ini'

        self.data_storage_path = tk.StringVar()
        self.data_storage_locked = False
        self.load_data_storage_path()

        self.station_contact_map = {
            "江宁维修站": "杨桂华",
            "禄口维修站": "李小康",
            "雨花维修站": "毛磊",
            "河西维修站": "吴月鹏",
            "城南维修站": "王大鹏",
            "城北维修站": "王正茂",
            "城东维修站": "高永金",
            "浦口维修站": "王创创",
            "江浦维修站": "王升超",
            "红山维修站": "孔维丹",
            "六合维修站": "周勇",
            "镇江西维修站": "陈伟",
            "句容恒大站": "戴群",
            "扬州东站": "叶鹏诚",
            "扬州西站": "王震磊"
        }

        self.maintenance_mode_var = tb.StringVar()

        self.create_widgets()

    def create_widgets(self):
        self.frame = tb.Frame(self.parent)
        self.frame.pack(fill='both', expand=True)

        self.path1 = tb.StringVar()
        self.path2 = tb.StringVar()
        self.output_path = tb.StringVar()
        self.text = tb.StringVar()
        self.maintenance_option = tb.StringVar(value="自保养")

        self.project_manager = tb.StringVar()
        self.contact1_selection = tb.StringVar()
        self.contact2_selection = tb.StringVar()

        self.project_contact = tb.StringVar()
        self.project_contact_phone = tb.StringVar()

        for i in range(25):
            self.frame.grid_rowconfigure(i, weight=1)
        for i in range(6):
            self.frame.grid_columnconfigure(i, weight=1)

        tb.Label(self.frame, text="Excel路径:").grid(row=0, column=1, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.path1, width=50).grid(row=0, column=2, sticky='we', padx=5, pady=5)
        tb.Button(self.frame, text="选择表格", command=self.select_excel, bootstyle="primary").grid(row=0, column=3, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="Word路径:").grid(row=1, column=1, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.path2, width=50).grid(row=1, column=2, sticky='we', padx=5, pady=5)
        tb.Button(self.frame, text="选择模版", command=self.select_word, bootstyle="primary").grid(row=1, column=3, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="读取的保养方式:").grid(row=1, column=4, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.maintenance_mode_var, width=20, state="readonly").grid(row=1, column=5, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="输出目录:").grid(row=2, column=1, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.output_path, width=50).grid(row=2, column=2, sticky='we', padx=5, pady=5)
        tb.Button(self.frame, text="选择输出目录", command=self.select_output_dir, bootstyle="primary").grid(row=2, column=3, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="请选择保养类型:").grid(row=3, column=1, sticky='e', padx=5, pady=5)
        tb.Radiobutton(self.frame, text="自保养", variable=self.maintenance_option, value="自保养").grid(row=3, column=2, sticky='w', padx=5, pady=5)
        tb.Radiobutton(self.frame, text="三方保养", variable=self.maintenance_option, value="三方保养").grid(row=3, column=3, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="项目经理:").grid(row=4, column=1, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.project_manager, width=30).grid(row=4, column=2, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="联系人1:").grid(row=5, column=1, sticky='e', padx=5, pady=5)
        contact1_combobox = tb.Combobox(
            self.frame,
            textvariable=self.contact1_selection,
            values=["不选择"] + list(self.station_contact_map.keys()),
            state="readonly",
            width=28
        )
        contact1_combobox.grid(row=5, column=2, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="联系人2:").grid(row=6, column=1, sticky='e', padx=5, pady=5)
        contact2_combobox = tb.Combobox(
            self.frame,
            textvariable=self.contact2_selection,
            values=["不选择", "刘峰", "徐铭崎", "赵凌", "杨磊侦"],
            state="readonly",
            width=28
        )
        contact2_combobox.grid(row=6, column=2, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="项目联系人:").grid(row=7, column=1, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.project_contact, width=30).grid(row=7, column=2, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="项目联系方式:").grid(row=8, column=1, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.project_contact_phone, width=30).grid(row=8, column=2, sticky='w', padx=5, pady=5)

        tb.Label(self.frame, text="数据存储位置:").grid(row=9, column=1, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.data_storage_path, width=50, state='readonly').grid(row=9, column=2, sticky='we', padx=5, pady=5)
        self.select_data_storage_button = tb.Button(self.frame, text="选择数据存储位置", command=self.select_data_storage_location, bootstyle="primary")
        self.select_data_storage_button.grid(row=9, column=3, sticky='w', padx=5, pady=5)

        self.modify_data_storage_button = tb.Button(self.frame, text="修改数据存储位置", command=self.modify_data_storage_location, bootstyle="secondary")
        self.modify_data_storage_button.grid(row=10, column=3, sticky='w', padx=5, pady=5)

        self.progress_var = tb.DoubleVar()
        self.progress_bar = tb.Progressbar(self.frame, variable=self.progress_var, maximum=100, bootstyle="info-striped")
        self.progress_bar.grid(row=11, column=1, columnspan=3, padx=10, pady=10, sticky='we')

        tb.Label(self.frame, textvariable=self.text).grid(row=12, column=1, columnspan=2, sticky='we', padx=5, pady=5)
        tb.Button(self.frame, text="开始运行", command=self.start_processing, bootstyle="success").grid(row=12, column=3, sticky='w', padx=5, pady=5)

        if self.data_storage_locked:
            self.select_data_storage_button.config(state='disabled')

    def load_data_storage_path(self):
        if os.path.exists(self.config_file):
            self.config.read(self.config_file)
            if 'Settings' in self.config and 'data_storage_path' in self.config['Settings']:
                data_storage = self.config['Settings']['data_storage_path']
                self.data_storage_path.set(data_storage)
                self.data_storage_locked = True
            else:
                self.data_storage_locked = False
        else:
            self.data_storage_locked = False

    def select_excel(self):
        file_path = filedialog.askopenfilename(filetypes=[("Excel 文件", "*.xlsx")])
        if file_path:
            self.path1.set(file_path)
            try:
                df_temp = pd.read_excel(file_path, sheet_name=0)
                if "预留保养方式" in df_temp.columns:
                    mode_value = str(df_temp["预留保养方式"].iloc[0]).strip()
                    self.maintenance_mode_var.set(mode_value)
                else:
                    self.maintenance_mode_var.set("（无此列）")
            except Exception as ex:
                messagebox.showerror("错误", f"读取 Excel 出错：{ex}")
                logging.error(f"读取 Excel 出错: {ex}")
                self.maintenance_mode_var.set("读取出错")

    def select_word(self):
        file_path = filedialog.askopenfilename(filetypes=[("Word 文件", "*.docx")])
        if file_path:
            self.path2.set(file_path)

    def select_output_dir(self):
        directory_path = filedialog.askdirectory()
        if directory_path:
            self.output_path.set(directory_path)

    def select_data_storage_location(self):
        if not self.data_storage_locked:
            directory_path = filedialog.askdirectory()
            if directory_path:
                self.data_storage_path.set(directory_path)
                self.data_storage_locked = True
                self.select_data_storage_button.config(state='disabled')
                if 'Settings' not in self.config:
                    self.config['Settings'] = {}
                self.config['Settings']['data_storage_path'] = directory_path
                with open(self.config_file, 'w') as configfile:
                    self.config.write(configfile)
        else:
            messagebox.showinfo("提示", "数据存储位置已锁定，无法修改。")

    def modify_data_storage_location(self):
        password = simpledialog.askstring("输入密码", "请输入密码：", show='*')
        if password == 'smec':
            self.data_storage_locked = False
            self.select_data_storage_button.config(state='normal')
            messagebox.showinfo("提示", "您可以重新选择数据存储位置。")
            self.select_data_storage_location()
        else:
            messagebox.showerror("错误", "密码错误，无法修改数据存储位置。")

    def start_processing(self):
        if not self.data_storage_path.get():
            messagebox.showerror("错误", "请先选择数据存储位置。")
            return
        self.frame.after(0, self.disable_widgets)
        threading.Thread(target=self.init).start()

    def disable_widgets(self):
        for widget in self.frame.winfo_children():
            if isinstance(widget, tb.Button):
                widget.config(state='disabled')

    def enable_widgets(self):
        for widget in self.frame.winfo_children():
            if isinstance(widget, tb.Button):
                widget.config(state='normal')
        if self.data_storage_locked:
            self.select_data_storage_button.config(state='disabled')

    def replace_text_in_word(self, file_name, old_text, new_text):
        docx = Document(file_name)
        for paragraph in docx.paragraphs:
            if old_text in paragraph.text:
                inline = paragraph.runs
                for i in range(len(inline)):
                    if old_text in inline[i].text:
                        text = inline[i].text.replace(str(old_text), str(new_text))
                        inline[i].text = text
                        inline[i].underline = WD_UNDERLINE.SINGLE

        for table in docx.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if old_text in paragraph.text:
                            inline = paragraph.runs
                            for i in range(len(inline)):
                                if old_text in inline[i].text:
                                    text = inline[i].text.replace(str(old_text), str(new_text))
                                    inline[i].text = text
                                    inline[i].underline = WD_UNDERLINE.SINGLE
        docx.save(file_name)

    def add_hash(self, x):
        return str(x) + '#'

    def cut_hash(self, x):
        return str(x).split("m/s")[0]

    def init(self):
        conn = None
        try:
            io_excel = self.path1.get()
            io_word = self.path2.get()
            output_directory = self.output_path.get()
            maintenance_type = self.maintenance_option.get()

            project_manager = self.project_manager.get()
            contact1_sel = self.contact1_selection.get()
            contact2_sel = self.contact2_selection.get()

            project_contact = self.project_contact.get()
            project_contact_phone = self.project_contact_phone.get()

            if not os.path.exists(io_excel):
                self.text.set("Excel 文件不存在！")
                logging.error("Excel 文件不存在")
                self.enable_widgets()
                return
            if not os.path.exists(io_word):
                self.text.set("Word 模板文件不存在！")
                logging.error("Word 模板文件不存在")
                self.enable_widgets()
                return
            if not os.path.exists(output_directory):
                self.text.set("输出目录不存在！")
                logging.error("输出目录不存在")
                self.enable_widgets()
                return

            if not project_manager:
                self.text.set("项目经理未填写！")
                logging.error("项目经理未填写")
                self.enable_widgets()
                return

            if not project_contact:
                self.text.set("项目联系人未填写！")
                logging.error("项目联系人未填写")
                self.enable_widgets()
                return
            if not project_contact_phone:
                self.text.set("项目联系方式未填写！")
                logging.error("项目联系方式未填写")
                self.enable_widgets()
                return

            self.text.set("正在处理，请稍候...")
            logging.info("开始处理")

            total_steps = 100
            for i in range(1, total_steps + 1):
                time.sleep(0.01)
                self.progress_var.set(i)
                self.frame.update_idletasks()

            df = pd.read_excel(io_excel, sheet_name=0, header=0,
                               usecols=["合同号-梯号", "速度/宽度", "产品型号", "层站门/提升高度", "预留保养期限", "使用方名称", "安装所在地"])
            df['预留保养期限'] = pd.to_numeric(df['预留保养期限'], errors='coerce').fillna(0).astype(int)

            user_name = df["使用方名称"].iloc[0]
            user_site = df["安装所在地"].iloc[0]
            user_date = df["预留保养期限"].iloc[0]

            df[["bianhao1", "bianhao2", '梯号']] = df['合同号-梯号'].str.split('-', expand=True)
            df['A'] = range(1, len(df) + 1)
            df['B'] = ''
            df['C'] = str(1)
            df = df[["A", "合同号-梯号", '梯号', "产品型号", "速度/宽度", "层站门/提升高度", "B", "C", "预留保养期限"]]
            df["梯号"] = df["梯号"].apply(self.add_hash)
            df["速度/宽度"] = df["速度/宽度"].apply(self.cut_hash)
            df = df.astype(str)

            doc = Document(io_word)
            table = doc.tables[1]
            current_rows = len(table.rows)
            if current_rows <= len(df):
                for _ in range(len(df) - current_rows + 1):
                    table.add_row()

            for i in range(len(df)):
                row_list = df.iloc[i].tolist()
                for e, item in enumerate(row_list):
                    table.cell(i + 1, e).text = item
                table.cell(i + 1, 6).text = "/"
                table.cell(i + 1, 9).text = "/"

            file_name = os.path.join(output_directory, f"{user_name}{maintenance_type}合同.docx")
            doc.save(file_name)

            if contact1_sel == "不选择":
                contact1_name = ""
            else:
                contact1_name = self.station_contact_map.get(contact1_sel, "")

            if contact2_sel == "不选择":
                contact2_name = ""
            else:
                contact2_name = contact2_sel

            contacts_combined = f"{contact1_name} {contact2_name}".strip()

            self.replace_text_in_word(file_name, "<<安装所在地>>", user_site)
            self.replace_text_in_word(file_name, "<<使用方名称>>", user_name)
            self.replace_text_in_word(file_name, "<<台数>>", str(len(df)))
            self.replace_text_in_word(file_name, "<<预留保养期限>>", str(user_date))
            self.replace_text_in_word(file_name, "<<项目经理>>", project_manager)
            self.replace_text_in_word(file_name, "<<联系人>>", contacts_combined)

            self.text.set(f'已完成，本次输入 {len(df)} 行数据！生成文件：{file_name}')
            logging.info(f'处理完成，生成文件：{file_name}')

            conn, cursor = self.setup_database()
            self.save_to_database(
                conn,
                cursor,
                project_manager,
                contact1_sel,
                contact2_sel,
                user_site,
                user_name,
                len(df),
                user_date,
                df['合同号-梯号'].tolist(),
                project_contact,
                project_contact_phone
            )

            self.export_database_to_excel(conn, cursor)

        except Exception as e:
            logging.error(f"处理过程中出错: {e}")
            self.show_error(traceback.format_exc())
        finally:
            self.enable_widgets()
            self.progress_var.set(0)
            if conn:
                conn.close()

    def setup_database(self):
        current_year = dt.datetime.now().year
        db_filename = f'免保合同制作记录_{current_year}.db'
        db_file_path = os.path.join(self.data_storage_path.get(), db_filename)
        os.makedirs(os.path.dirname(db_file_path), exist_ok=True)

        conn = sqlite3.connect(db_file_path)
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                time TEXT,
                project_manager TEXT,
                contact1 TEXT,
                contact2 TEXT,
                install_location TEXT,
                user_name TEXT,
                num_units INTEGER,
                maintenance_period INTEGER,
                contract_numbers TEXT,
                project_contact TEXT,
                project_contact_phone TEXT
            )
        ''')
        conn.commit()
        return conn, cursor

    def save_to_database(self, conn, cursor,
                         project_manager, contact1, contact2,
                         install_location, user_name,
                         num_units, maintenance_period,
                         contract_numbers,
                         project_contact, project_contact_phone):
        try:
            current_time = dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            contract_numbers_str = ','.join(contract_numbers)

            if isinstance(maintenance_period, bytes):
                maintenance_period = int.from_bytes(maintenance_period, byteorder='little')
            else:
                maintenance_period = int(maintenance_period)

            cursor.execute('''
                INSERT INTO records (
                    time, project_manager, contact1, contact2,
                    install_location, user_name, num_units,
                    maintenance_period, contract_numbers,
                    project_contact, project_contact_phone
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                current_time, project_manager, contact1, contact2,
                install_location, user_name, num_units,
                maintenance_period, contract_numbers_str,
                project_contact, project_contact_phone
            ))
            conn.commit()
            logging.info("数据已成功保存到本地 SQLite 数据库。")
        except Exception as e:
            logging.error(f"保存到 SQLite 数据库时出错: {e}")
            self.show_error(f"保存到 SQLite 数据库时出错: {e}")

    def export_database_to_excel(self, conn, cursor):
        try:
            cursor.execute('SELECT * FROM records')
            rows = cursor.fetchall()
            column_names = [desc[0] for desc in cursor.description]

            df_db = pd.DataFrame(rows, columns=column_names)

            if 'maintenance_period' in df_db.columns:
                df_db['maintenance_period'] = df_db['maintenance_period'].apply(
                    lambda x: int.from_bytes(x, byteorder='little') if isinstance(x, bytes) else x
                )

            column_mapping = {
                'id': '编号',
                'time': '时间',
                'project_manager': '项目经理',
                'contact1': '联系人1',
                'contact2': '联系人2',
                'install_location': '安装所在地',
                'user_name': '使用方名称',
                'num_units': '台数',
                'maintenance_period': '预留保养期限',
                'contract_numbers': '合同号-梯号',
                'project_contact': '项目联系人',
                'project_contact_phone': '项目联系方式'
            }
            df_db.rename(columns=column_mapping, inplace=True)

            excel_file = os.path.join(self.data_storage_path.get(), '免保合同制作记录.xlsx')
            if os.path.exists(excel_file):
                df_excel = pd.read_excel(excel_file)
                df_combined = pd.concat([df_excel, df_db], ignore_index=True)
                df_combined.drop_duplicates(subset=df_combined.columns, keep='first', inplace=True)
                df_combined.to_excel(excel_file, index=False)
            else:
                df_db.to_excel(excel_file, index=False)

            logging.info("数据已成功导出到 Excel 文件。")
        except Exception as e:
            logging.error(f"导出数据到 Excel 文件时出错: {e}")
            self.show_error(f"导出数据到 Excel 文件时出错: {e}")

    def show_error(self, error_text):
        error_window = tb.Toplevel(self.parent)
        error_window.title("报错窗口")
        error_window.geometry("800x500")
        text_widget = tb.Text(error_window, height=30, width=80)
        text_widget.insert('insert', error_text)
        text_widget.pack()
        logging.error(f"错误信息：\n{error_text}")


class WordExcelProcessorApp:
    """
    终极批量合同工具：直接在界面读取Excel名称预览，并由界面输入金额/期限等参数，
    不再需要在Excel中修改Config后再制作。
    """

    def __init__(self, parent):
        self.parent = parent
        self.frame = tb.Frame(self.parent)
        self.frame.pack(fill='both', expand=True)

        self.text = tb.StringVar()
        self.progress_var = tb.DoubleVar()
        self.output_path = tb.StringVar()
        self.template_folder = tb.StringVar()
        self.excel_folder = tb.StringVar()

        self.word_template_choice_var = tb.StringVar(value="老版自保养合同模版A")
        self.template_name_choices = [
            "老版自保养合同模版A",
            "老版自保养合同模版B",
            "老版自保养合同模版C",
            "新版自保养合同模版A1",
            "新版自保养合同模版A2",
            "新版自保养合同模版B1",
            "新版自保养合同模版B2",
            "新版自保养合同模版B3",
            "新版自保养合同模版B4",
            "新版自保养合同模版C",
        ]

        self.start_date_var = tb.StringVar()
        self.end_date_var = tb.StringVar()
        self.pay_method_var = tb.StringVar(value="季度付")
        self.annual_price_var = tb.StringVar(value="12000")
        self.subtotal_var = tb.StringVar(value="3000")
        self.station_choice_var = tb.StringVar(value="江宁站")
        self.contact1_var = tb.StringVar(value="张勇辉")
        self.contact2_var = tb.StringVar(value="徐铭崎")
        self.annual_inspection_var = tb.StringVar(value="不包含")

        self.loaded_excels = []

        self.create_widgets()

    def create_widgets(self):
        tb.Label(self.frame, text="Excel文件夹:").grid(row=0, column=0, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.excel_folder, width=45).grid(row=0, column=1, padx=5, pady=5, sticky='we')
        tb.Button(self.frame, text="加载Excel文件夹", command=self.load_excel_folder, bootstyle="primary").grid(row=0, column=2, padx=5, pady=5)

        tb.Label(self.frame, text="模板文件夹:").grid(row=1, column=0, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.template_folder, width=45).grid(row=1, column=1, padx=5, pady=5, sticky='we')
        tb.Button(self.frame, text="选择模板文件夹", command=self.select_template_folder).grid(row=1, column=2, padx=5, pady=5)

        tb.Label(self.frame, text="输出目录:").grid(row=2, column=0, sticky='e', padx=5, pady=5)
        tb.Entry(self.frame, textvariable=self.output_path, width=45).grid(row=2, column=1, padx=5, pady=5, sticky='we')
        tb.Button(self.frame, text="选择输出目录", command=self.select_output_dir).grid(row=2, column=2, padx=5, pady=5)

        tb.Label(self.frame, text="Word模板:").grid(row=3, column=0, sticky='e', padx=5, pady=5)
        tb.Combobox(
            self.frame,
            textvariable=self.word_template_choice_var,
            values=self.template_name_choices,
            state="readonly",
            width=42
        ).grid(row=3, column=1, padx=5, pady=5, sticky='w')

        tb.Label(self.frame, text="维修站(联系人1)/联系人2:").grid(row=4, column=0, sticky='e', padx=5, pady=5)
        contact_frame = tb.Frame(self.frame)
        contact_frame.grid(row=4, column=1, sticky='w', padx=5, pady=5)
        station_box = tb.Combobox(
            contact_frame,
            textvariable=self.station_choice_var,
            values=["江宁站", "禄口站"],
            state="readonly",
            width=18
        )
        station_box.grid(row=0, column=0, padx=5)
        station_box.bind('<<ComboboxSelected>>', lambda e: self.update_contacts_by_station())
        tb.Entry(contact_frame, textvariable=self.contact1_var, width=15, state="readonly").grid(row=0, column=1, padx=5)
        tb.Entry(contact_frame, textvariable=self.contact2_var, width=15, state="readonly").grid(row=0, column=2, padx=5)

        tb.Label(self.frame, text="开始日期/结束日期:").grid(row=5, column=0, sticky='e', padx=5, pady=5)
        date_frame = tb.Frame(self.frame)
        date_frame.grid(row=5, column=1, sticky='w', padx=5, pady=5)
        tb.Entry(date_frame, textvariable=self.start_date_var, width=18).grid(row=0, column=0, padx=5)
        tb.Entry(date_frame, textvariable=self.end_date_var, width=18).grid(row=0, column=1, padx=5)
        tb.Label(date_frame, text="示例：2025年3月30日").grid(row=0, column=2, padx=5)

        tb.Label(self.frame, text="付款方式:").grid(row=6, column=0, sticky='e', padx=5, pady=5)
        tb.Combobox(
            self.frame,
            textvariable=self.pay_method_var,
            values=["季度付", "半年付", "年付"],
            state="readonly",
            width=12
        ).grid(row=6, column=1, padx=5, pady=5, sticky='w')

        tb.Label(self.frame, text="年价/小计:").grid(row=7, column=0, sticky='e', padx=5, pady=5)
        price_frame = tb.Frame(self.frame)
        price_frame.grid(row=7, column=1, sticky='w', padx=5, pady=5)
        tb.Entry(price_frame, textvariable=self.annual_price_var, width=15).grid(row=0, column=0, padx=5)
        tb.Entry(price_frame, textvariable=self.subtotal_var, width=15).grid(row=0, column=1, padx=5)

        tb.Label(self.frame, text="年检费:").grid(row=8, column=0, sticky='e', padx=5, pady=5)
        tb.Combobox(
            self.frame,
            textvariable=self.annual_inspection_var,
            values=["包含", "不包含"],
            state="readonly",
            width=10
        ).grid(row=8, column=1, padx=5, pady=5, sticky='w')

        tb.Button(
            self.frame,
            text="批量生成合同",
            command=self.start_batch_generate_contracts,
            bootstyle="success"
        ).grid(row=9, column=0, columnspan=3, padx=10, pady=10, sticky='we')

        self.progress_bar = tb.Progressbar(self.frame, variable=self.progress_var, maximum=100, bootstyle="info-striped")
        self.progress_bar.grid(row=10, column=0, columnspan=3, padx=10, pady=10, sticky='we')
        tb.Label(self.frame, textvariable=self.text).grid(row=11, column=0, columnspan=3, padx=5, pady=5, sticky='w')

        columns = ("filename", "project", "address", "client", "rows")
        self.preview = tb.Treeview(self.frame, columns=columns, show="headings", height=10)
        headings = {
            "filename": "文件名",
            "project": "项目名称",
            "address": "项目地址",
            "client": "保养客户名称",
            "rows": "台数"
        }
        for col, text in headings.items():
            self.preview.heading(col, text=text)
            self.preview.column(col, width=140, anchor='center')
        self.preview.grid(row=12, column=0, columnspan=3, padx=10, pady=10, sticky='nsew')

        self.frame.grid_columnconfigure(1, weight=1)
        self.frame.grid_rowconfigure(12, weight=1)
        self.update_contacts_by_station()

    def select_template_folder(self):
        folder = filedialog.askdirectory(title="选择模板文件夹(docx)")
        if folder:
            self.template_folder.set(folder)

    def select_output_dir(self):
        folder = filedialog.askdirectory(title="选择输出目录(生成合同)")
        if folder:
            self.output_path.set(folder)

    def load_excel_folder(self):
        folder = filedialog.askdirectory(title="选择Excel文件夹")
        if not folder:
            return
        self.excel_folder.set(folder)
        self.loaded_excels = self.read_excel_preview(folder)
        self.refresh_preview()
        self.text.set(f"已加载 {len(self.loaded_excels)} 个Excel")

    def read_excel_preview(self, folder):
        excels = []
        for fname in os.listdir(folder):
            if not fname.endswith('.xlsx'):
                continue
            path = os.path.join(folder, fname)
            try:
                df_main = pd.read_excel(path, sheet_name=0)
                project_name = str(df_main.get("项目名称", pd.Series([""])).iloc[0]).strip()
                project_addr = str(df_main.get("项目地址", pd.Series([""])).iloc[0]).strip()
                client_name = str(df_main.get("保养客户名称", pd.Series([""])).iloc[0]).strip()
                row_count = len(df_main.index)
            except Exception:
                project_name = "读取失败"
                project_addr = "读取失败"
                client_name = "读取失败"
                row_count = 0
            excels.append({
                "filename": fname,
                "path": path,
                "project": project_name,
                "address": project_addr,
                "client": client_name,
                "rows": row_count
            })
        return excels

    def refresh_preview(self):
        for row in self.preview.get_children():
            self.preview.delete(row)
        for item in self.loaded_excels:
            self.preview.insert('', 'end', values=(
                item["filename"], item["project"], item["address"], item["client"], item["rows"]
            ))

    def update_contacts_by_station(self):
        station = self.station_choice_var.get()
        if station == "禄口站":
            self.contact1_var.set("李小康")
        else:
            self.contact1_var.set("张勇辉")
        self.contact2_var.set("徐铭崎")

    def start_batch_generate_contracts(self):
        if not self.loaded_excels:
            self.text.set("请先加载Excel文件夹")
            return
        if not self.template_folder.get() or not self.output_path.get():
            self.text.set("请先选择模板文件夹和输出目录")
            return
        threading.Thread(target=self.batch_generate_contracts_thread).start()

    def batch_generate_contracts_thread(self):
        try:
            self.disable_widgets()
            total = len(self.loaded_excels)
            self.text.set("开始批量生成合同...")
            for idx, info in enumerate(self.loaded_excels, start=1):
                self.progress_var.set(idx / total * 100)
                self.frame.update_idletasks()
                self.text.set(f"({idx}/{total}) 正在处理: {info['filename']}")
                try:
                    self.generate_contract_for_single(info)
                except Exception as exc:
                    logging.error(f"生成合同失败 {info['filename']}: {exc}")
                    self.show_error(traceback.format_exc())
            self.text.set("批量生成合同完成")
        except Exception:
            self.show_error(traceback.format_exc())
        finally:
            self.progress_var.set(0)
            self.enable_widgets()

    def generate_contract_for_single(self, excel_info):
        excel_path = excel_info["path"]
        self.update_contacts_by_station()
        start_dt = self.parse_chinese_date(self.start_date_var.get().strip())
        end_dt = self.parse_chinese_date(self.end_date_var.get().strip())
        pay_method = self.pay_method_var.get().strip()
        annual_price_str = self.annual_price_var.get().strip()
        subtotal_str = self.subtotal_var.get().strip()
        contact1 = self.contact1_var.get().strip()
        contact2 = self.contact2_var.get().strip()
        word_tpl = self.word_template_choice_var.get().strip()
        annual_inspection = self.annual_inspection_var.get().strip()

        if not start_dt or not end_dt:
            raise ValueError("请填写有效的开始日期和结束日期（例：2025年3月30日）")

        service_months = 0
        if end_dt >= start_dt:
            delta = relativedelta(end_dt + timedelta(days=1), start_dt)
            service_months = delta.years * 12 + delta.months
        service_months_str = str(service_months)

        df_main = pd.read_excel(excel_path, sheet_name=0, header=0,
                                usecols=["合同号-梯号", "速度/宽度", "产品型号", "层站门/提升高度", "项目名称", "项目地址", "保养客户名称"])
        if df_main.empty:
            raise ValueError(f"主数据为空: {excel_path}")

        user_name = str(df_main["项目名称"].iloc[0])
        user_site = str(df_main["项目地址"].iloc[0])
        user_ming = str(df_main["保养客户名称"].iloc[0])

        df_main[['bianhao1', 'bianhao2', '梯号']] = df_main["合同号-梯号"].str.split('-', expand=True)
        df_main['A'] = range(1, len(df_main) + 1)
        df_main['B'] = ''
        df_main['C'] = '1'
        df_main = df_main[["A", "合同号-梯号", "梯号", "产品型号", "速度/宽度", "层站门/提升高度", "B", "C"]]
        df_main = df_main.astype(str)

        docx_path = os.path.join(self.template_folder.get(), word_tpl + ".docx")
        if not os.path.exists(docx_path):
            raise FileNotFoundError(f"找不到Word模板: {docx_path}")

        doc = Document(docx_path)
        if len(doc.tables) < 2:
            raise ValueError("Word模板不足2个表，无法写第二表")
        table2 = doc.tables[1]
        current_rows = len(table2.rows)
        required_rows = len(df_main) + 1
        if current_rows < required_rows:
            for _ in range(required_rows - current_rows):
                table2.add_row()

        for i in range(len(df_main)):
            row_data = df_main.iloc[i].tolist()
            for e, val in enumerate(row_data):
                if e == 4 and isinstance(val, str):
                    val = val.replace("m/s", "").strip()
                cell = table2.cell(i + 1, e)
                cell.text = str(val)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for col_index, value in zip([6, 8, 9], [annual_price_str, service_months_str, subtotal_str]):
                cell = table2.cell(i + 1, col_index)
                cell.text = str(value)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if annual_inspection == "包含":
            self.insert_paragraph_after_table(table2, "备注：该合同包含电梯年检费用", alignment=WD_ALIGN_PARAGRAPH.LEFT)

        count_rows = len(df_main)
        try:
            sub_val = float(subtotal_str)
            total_amount = int(sub_val * count_rows)
        except Exception:
            total_amount = 0
        total_amount_str = str(total_amount)
        total_amount_upper = self.number_to_chinese_upper(total_amount)

        if len(doc.tables) >= 4:
            table4 = doc.tables[3]
            self.fill_payment_table(table4, pay_method, total_amount, start_dt, end_dt)
        else:
            raise ValueError("Word模板不足4个表，无法写付款计划")

        out_name = f"{user_ming}自保养合同.docx"
        if start_dt.year == 2025:
            out_name = f"{user_ming}2025自保养合同.docx"
        out_path = os.path.join(self.output_path.get(), out_name)
        doc.save(out_path)

        self.replace_placeholder(out_path, "<<项目名称>>", user_name)
        self.replace_placeholder(out_path, "<<项目地址>>", user_site)
        self.replace_placeholder(out_path, "<<开始保养时间>>", start_dt.strftime("%Y年%m月%d日"))
        self.replace_placeholder(out_path, "<<结束保养时间>>", end_dt.strftime("%Y年%m月%d日"))
        self.replace_placeholder(out_path, "<<总金额>>", total_amount_str)
        self.replace_placeholder(out_path, "<<大写总金额>>", total_amount_upper)
        self.replace_placeholder(out_path, "<<台数>>", str(count_rows))
        contacts = f"{contact1}   {contact2}".strip()
        self.replace_placeholder(out_path, "<<联系人>>", contacts)

        logging.info(f"生成: {out_path}")

    def fill_payment_table(self, table4, pay_method, total_amount, start_dt, end_dt):
        schedule = []
        if pay_method == "季度付":
            installments = 4
            each = total_amount / installments if installments else 0
            for i in range(1, installments + 1):
                due_date = start_dt + relativedelta(months=3 * i) - timedelta(days=1)
                due_date_str = f"{due_date.strftime('%Y年%m月%d日')}前"
                schedule.append((i, due_date_str, each))
            for _ in range(installments + 1, 9):
                schedule.append((None, None, None))
        elif pay_method == "半年付":
            installments = 2
            each = total_amount / installments if installments else 0
            for i in range(1, installments + 1):
                due_date = start_dt + relativedelta(months=6 * i) - timedelta(days=1)
                due_date_str = f"{due_date.strftime('%Y年%m月%d日')}前"
                schedule.append((i, due_date_str, each))
            for _ in range(installments + 1, 9):
                schedule.append((None, None, None))
        else:
            installments = 1
            due_date = end_dt if end_dt else start_dt + relativedelta(years=1) - timedelta(days=1)
            due_date_str = f"{due_date.strftime('%Y年%m月%d日')}前"
            schedule.append((1, due_date_str, total_amount))
            for _ in range(2, 9):
                schedule.append((None, None, None))

        if len(table4.rows) < 5:
            raise ValueError("第4表行数不足(需≥5行)，无法写付款计划")

        for row_idx in range(1, 5):
            row = table4.rows[row_idx]
            left_seq = row_idx
            right_seq = row_idx + 4
            left_item = schedule[left_seq - 1]
            right_item = schedule[right_seq - 1]

            row.cells[0].text = str(left_seq)
            seqL, dtL, amtL = left_item
            if seqL is None:
                row.cells[1].text = "/年 /月 /日前"
                row.cells[2].text = "/"
            else:
                row.cells[1].text = dtL
                row.cells[2].text = f"{amtL:.2f}"

            row.cells[3].text = str(right_seq)
            seqR, dtR, amtR = right_item
            if seqR is None:
                row.cells[4].text = "/年 /月 /日前"
                row.cells[5].text = "/"
            else:
                row.cells[4].text = dtR
                row.cells[5].text = f"{amtR:.2f}"

        for row_idx in range(1, 5):
            row = table4.rows[row_idx]
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                    paragraph.paragraph_format.line_spacing = 1.5

    def replace_placeholder(self, file_name, old_text, new_text):
        docx = Document(file_name)
        for paragraph in docx.paragraphs:
            if old_text in paragraph.text:
                for run in paragraph.runs:
                    if old_text in run.text:
                        run.text = run.text.replace(old_text, new_text)
                        run.underline = WD_UNDERLINE.SINGLE

        for table in docx.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if old_text in paragraph.text:
                            for run in paragraph.runs:
                                if old_text in run.text:
                                    run.text = run.text.replace(old_text, new_text)
                                    run.underline = WD_UNDERLINE.SINGLE
        docx.save(file_name)

    def parse_chinese_date(self, date_str):
        if not date_str:
            return None
        patterns = [
            r'(\d{4})年(\d{1,2})月(\d{1,2})日',
            r'(\d{4})-(\d{1,2})-(\d{1,2})'
        ]
        for pat in patterns:
            m = re.match(pat, date_str)
            if m:
                y, mo, dy = m.groups()
                return dt.datetime(int(y), int(mo), int(dy))
        return None

    def number_to_chinese_upper(self, num):
        digits = ['零', '壹', '贰', '叁', '肆', '伍', '陆', '柒', '捌', '玖']
        units = ['', '拾', '佰', '仟']
        higher_units = ['', '万', '亿', '兆']
        s = str(num)
        length = len(s)
        result = ''
        zero = False
        for i, digit in enumerate(s):
            n = int(digit)
            pos = length - i - 1
            unit_pos = pos % 4
            higher_pos = pos // 4
            if n == 0:
                zero = True
            else:
                if zero:
                    result += '零'
                    zero = False
                result += digits[n] + units[unit_pos]
            if unit_pos == 0 and pos != 0:
                result += higher_units[higher_pos]
        result = result.replace("壹拾", "拾")
        while '零零' in result:
            result = result.replace('零零', '零')
        result = result.rstrip('零')
        return result if result else '零'

    def insert_paragraph_after_table(self, table, text, alignment=WD_ALIGN_PARAGRAPH.LEFT):
        tbl = table._tbl
        parent = tbl.getparent()
        new_p = OxmlElement("w:p")
        tbl.addnext(new_p)
        new_para = Paragraph(new_p, parent)
        new_para.text = text
        new_para.alignment = alignment

    def show_error(self, err_text):
        w = tb.Toplevel(self.parent)
        w.title("错误信息")
        w.geometry("800x400")
        txt = tb.Text(w, width=80, height=20)
        txt.insert('1.0', err_text)
        txt.pack()

    def disable_widgets(self):
        for widget in self.frame.winfo_children():
            if isinstance(widget, tb.Button):
                widget.config(state='disabled')

    def enable_widgets(self):
        for widget in self.frame.winfo_children():
            if isinstance(widget, tb.Button):
                widget.config(state='normal')


class MainApplication:
    def __init__(self, root):
        self.root = root
        self.style = tb.Style(theme='flatly')
        self.root.title("合同制作工具")
        self.root.geometry("1050x800")
        self.create_widgets()

    def create_widgets(self):
        self.left_frame = tb.Frame(self.root)
        self.left_frame.pack(side='left', fill='y')

        self.right_frame = tb.Frame(self.root)
        self.right_frame.pack(side='right', expand=True, fill='both')

        self.function_button = tb.Button(
            self.left_frame,
            text="免保合同制作",
            command=self.show_free_maintenance_contract,
            bootstyle="success"
        )
        self.function_button.pack(fill='x', padx=10, pady=10)

        self.batch_button = tb.Button(
            self.left_frame,
            text="批量合同工具",
            command=self.show_word_excel_processor,
            bootstyle="primary"
        )
        self.batch_button.pack(fill='x', padx=10, pady=10)

        self.current_function = None
        self.show_free_maintenance_contract()

    def show_free_maintenance_contract(self):
        self.clear_right_frame()
        self.current_function = FreeMaintenanceContractApp(self.right_frame)

    def show_word_excel_processor(self):
        self.clear_right_frame()
        self.current_function = WordExcelProcessorApp(self.right_frame)

    def clear_right_frame(self):
        for widget in self.right_frame.winfo_children():
            widget.destroy()


def main():
    root = tb.Window(themename="flatly")
    MainApplication(root)
    root.mainloop()


if __name__ == "__main__":
    main()
